# Download official BEA underlying-detail tables and build a PCE inflation workbook.
import argparse
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import requests

ROOT = Path(__file__).resolve().parent
URL = 'https://apps.bea.gov/national/Release/XLS/Underlying/Section2All_xls.xlsx'
parser = argparse.ArgumentParser()
parser.add_argument('--refresh', action='store_true', help='Download the latest complete BEA section, including revisions.')
args = parser.parse_args()
raw = ROOT / 'raw' / 'Section2All.xlsx'
raw.parent.mkdir(parents=True, exist_ok=True)
if args.refresh or not raw.exists():
    response = requests.get(URL, timeout=120)
    response.raise_for_status()
    if not response.content.startswith(b'PK'):
        raise ValueError('BEA returned something other than an XLSX file')
    candidate = raw.with_suffix('.download.xlsx')
    candidate.write_bytes(response.content)
    check = openpyxl.load_workbook(candidate, read_only=True)
    assert {'U20404-M', 'U20405-M'}.issubset(check.sheetnames)
    check.close()
    if raw.exists():
        old_hash = hashlib.sha256(raw.read_bytes()).hexdigest()[:12]
        archive = raw.parent / f'Section2All_{old_hash}.xlsx'
        if not archive.exists():
            archive.write_bytes(raw.read_bytes())
    candidate.replace(raw)

book = openpyxl.load_workbook(raw, read_only=True, data_only=True)


def read_table(sheet_name):
    rows = list(book[sheet_name].values)
    header = next(i for i, row in enumerate(rows) if row[0] == 'Line')
    periods = rows[header][3:]
    data = [r for r in rows[header+1:] if str(r[0]).isdigit() and r[2]]
    meta = pd.DataFrame([{'Line': int(r[0]), 'Name': r[1].strip(), 'Code': r[2],
                          'Indent': (len(r[1]) - len(r[1].lstrip())) // 2} for r in data]).set_index('Line')
    values = pd.DataFrame([r[3:] for r in data], index=meta.index, columns=periods)
    values = values.replace('.....', np.nan).apply(pd.to_numeric, errors='raise').T
    if sheet_name.endswith('-M'):
        values.index = pd.to_datetime(values.index, format='%YM%m')
    values.index.name = 'Period'
    return meta, values, [r[0] for r in rows[:header] if r[0]]


meta, prices, price_notes = read_table('U20404-M')
ex_meta, spending, expense_notes = read_table('U20405-M')
pd.testing.assert_index_equal(meta.index, ex_meta.index)
assert meta['Name'].equals(ex_meta['Name'])
assert prices.index.equals(pd.date_range(prices.index.min(), prices.index.max(), freq='MS'))
assert prices.columns.is_unique and spending.columns.is_unique
assert len(meta) == 402, 'BEA table changed: review hierarchy boundary before rebuilding'
meta = meta.rename(columns={'Code': 'Price code'})
meta['Spending code'] = ex_meta['Code']
meta['Parent line'] = pd.Series(dtype='Int64')
meta['Depth'] = pd.Series(dtype='Int64')
meta['Sign to parent'] = 1
meta['Sign to root'] = 1
meta['Basket member'] = meta.index <= 368
meta.loc[1, 'Depth'] = 0
stack = [1]
children = {i: [] for i in range(1, 369)}
for line in range(2, 369):
    depth = int(meta.loc[line, 'Indent']) + 1
    while int(meta.loc[stack[-1], 'Depth']) >= depth:
        stack.pop()
    parent = stack[-1]
    assert int(meta.loc[parent, 'Depth']) == depth - 1
    sign = -1 if meta.loc[line, 'Name'].startswith('Less:') else 1
    meta.loc[line, ['Parent line', 'Depth', 'Sign to parent', 'Sign to root']] = [parent, depth, sign, sign * meta.loc[parent, 'Sign to root']]
    children[parent].append(line)
    stack.append(line)

# Check the inferred hierarchy against independently published current-dollar totals.
audit = []
for parent, child in children.items():
    if not child:
        continue
    subtotal = spending[child].mul(meta.loc[child, 'Sign to parent'], axis=1).sum(axis=1, min_count=len(child))
    residual = subtotal - spending[parent]
    tolerance = (len(child) + 1) * 0.5 + 1e-8
    if residual.abs().max() > tolerance:
        raise ValueError(f'Hierarchy fails published rounding tolerance at line {parent}: {residual.abs().max()}')
    audit.append({'Parent line': parent, 'Name': meta.loc[parent, 'Name'], 'Children': len(child),
                  'Complete months': int(residual.notna().sum()), 'Maximum absolute residual USD million SAAR': float(residual.abs().max()),
                  'Rounding tolerance': tolerance})

# Reconcile published rounding recursively. Missing siblings prevent a finer allocation.
allocation = pd.DataFrame(index=spending.index, columns=range(1, 369), dtype=float)
allocation[1] = 1.0
for parent, child in children.items():
    if child:
        signed = spending[child].mul(meta.loc[child, 'Sign to parent'], axis=1)
        denominator = signed.sum(axis=1, min_count=len(child)).replace(0, np.nan)
        allocation[child] = signed.div(denominator, axis=0).mul(allocation[parent], axis=0)

latest = spending.index.max()
meta['Latest index'] = prices.loc[latest]
meta['Latest spending USD million SAAR'] = spending.loc[latest]
meta['Latest published spending share'] = spending.loc[latest] / spending.loc[latest, 1]
meta['Latest reconciled signed share'] = allocation.loc[latest]
meta['First observation'] = [prices[i].first_valid_index() for i in meta.index]
meta['Last observation'] = [prices[i].last_valid_index() for i in meta.index]
meta['Monthly observations'] = prices.notna().sum()
tiers, tier_audit = {}, []
cut = [1]
for depth in range(1, int(meta['Depth'].max()) + 1):
    # Lines 157–159 are alternative historical housing breakdowns, not simultaneous siblings.
    cut = [c for parent in cut for c in (children[parent] if children[parent] and parent != 156 else [parent])]
    weights = allocation[cut]
    sums = weights.sum(axis=1, min_count=len(cut))
    assert pd.notna(sums.loc[latest]), f'Latest tier {depth} has incomplete weights'
    error = (sums.dropna() - 1).abs().max()
    assert error < 1e-12
    tiers[depth] = list(cut)
    tier_audit.append({'Tier': depth, 'Components': len(cut), 'Latest sum': float(sums.loc[latest]),
                       'Complete weight months': int(sums.notna().sum()), 'Maximum sum error': float(error),
                       'Negative latest shares': int((weights.loc[latest] < 0).sum())})

mom = prices.pct_change(1, fill_method=None) * 100
yoy = prices.pct_change(12, fill_method=None) * 100
notes = [
    'US PCE inflation — official BEA underlying-detail data',
    f'Monthly coverage {prices.index.min():%Y-%m} to {latest:%Y-%m}; {len(meta)} table rows including overlapping additional aggregates.',
    *price_notes[1:6],
    'Monthly prices are seasonally adjusted, 2017=100. Spending is millions of current dollars, seasonally adjusted annual rates.',
    'All prices contains every Table 2.4.4U row. Column labels are BEA line numbers; Series supplies names and price/spending codes.',
    'Table line numbers align prices and spending. Some additional aggregates repeat series codes; those rows are retained.',
    'Each tier is an exhaustive cut of the BEA indentation hierarchy, carrying terminal items forward into deeper tiers.',
    'Only lines 1–368 form the basket tree; additional and market-based aggregates are excluded from tier totals.',
    'Line 156 is carried forward at the deepest tier: its lines 157–159 are alternative historical housing breakdowns, never a complete simultaneous split.',
    'Weights vary monthly. Published shares are current-dollar spending divided by total PCE; they approximate relative importance.',
    'Reconciled shares allocate each parent by signed sibling expenditure proportions, removing only published rounding discrepancies.',
    'Each fully observed tier sums to 1; incomplete historical tiers remain incomplete, with no filling or renormalisation over observed items.',
    'Deep tiers contain signed accounting deductions (Less:) and negative net spending, so some weights are negative. These are not positive-only baskets.',
    'Signs propagate to descendants of Less: lines. Original spending and price values are preserved unchanged.',
    'PCE is a Fisher chain-type index: these shares do NOT exactly reproduce the official index or its component contributions.',
    'MoM and YoY are unannualised percentage changes. Missing observations and undefined changes are blank.',
    'Annual and quarterly official price and spending histories are included separately, in their published period labels.',
    'Latest weights ranks are descending by signed share within each tier.',
    'Source: ' + URL,
    'Weights method: https://www.bea.gov/help/faq/1006',
    'Refresh: python build_pce.py --refresh downloads the current complete file and rebuilds all histories, including revisions.'
]
output = ROOT.parent / 'USA PCE - Historical Data and Weights.xlsx'
labels = {i: f'{i} | {meta.loc[i, "Name"]}' for i in meta.index}
latest_weights = []
for tier, lines in tiers.items():
    frame = meta.loc[lines].copy()
    frame.insert(0, 'Tier', tier)
    frame['Weight rank'] = frame['Latest reconciled signed share'].rank(method='min', ascending=False).astype(int)
    latest_weights.append(frame.reset_index())

with pd.ExcelWriter(output, engine='xlsxwriter', datetime_format='yyyy-mm') as writer:
    pd.DataFrame({'Read Me': notes}).to_excel(writer, sheet_name='Read Me', index=False)
    meta.to_excel(writer, sheet_name='Series')
    prices.rename(columns=labels).to_excel(writer, sheet_name='All monthly prices')
    spending.rename(columns=labels).to_excel(writer, sheet_name='All monthly spending')
    mom.rename(columns=labels).to_excel(writer, sheet_name='MoM percent')
    yoy.rename(columns=labels).to_excel(writer, sheet_name='YoY percent')
    pd.concat(latest_weights, ignore_index=True).to_excel(writer, sheet_name='Latest weights and ranks', index=False)
    pd.DataFrame(tier_audit).to_excel(writer, sheet_name='Tier checks', index=False)
    pd.DataFrame(audit).to_excel(writer, sheet_name='Hierarchy checks', index=False)
    spending.div(spending[1], axis=0).rename(columns=labels).to_excel(writer, sheet_name='Published spending shares')
    for tier, lines in tiers.items():
        prices[lines].rename(columns=labels).to_excel(writer, sheet_name=f'Tier {tier} prices')
        frame = allocation[lines].rename(columns=labels).copy()
        frame['SUM (blank if incomplete)'] = allocation[lines].sum(axis=1, min_count=len(lines))
        frame.to_excel(writer, sheet_name=f'Tier {tier} weights')
    for frequency in ['A', 'Q']:
        for table, name in [('U20404', 'prices'), ('U20405', 'spending')]:
            _, values, _ = read_table(f'{table}-{frequency}')
            values.rename(columns=labels).to_excel(writer, sheet_name=f'{frequency} {name}')
    for name, sheet in writer.sheets.items():
        sheet.freeze_panes(1, 1)
        sheet.set_column(0, 0, 18)
        sheet.set_column(1, 420, 18)
        if name not in ['Read Me']:
            sheet.autofilter(0, 0, sheet.dim_rowmax, sheet.dim_colmax)
    writer.sheets['Read Me'].set_column(0, 0, 145)
    writer.sheets['Series'].set_column(1, 1, 75)
    writer.sheets['Latest weights and ranks'].set_column(2, 2, 75)
    writer.sheets['Read Me'].activate()

prices.to_csv(ROOT / 'monthly_prices.csv', date_format='%Y-%m')
spending.to_csv(ROOT / 'monthly_spending.csv', date_format='%Y-%m')
meta.to_csv(ROOT / 'series.csv')
manifest = {'source': URL, 'built_at_utc': datetime.now(timezone.utc).isoformat(),
            'raw_sha256': hashlib.sha256(raw.read_bytes()).hexdigest(), 'release_notes': price_notes,
            'first_month': str(prices.index.min().date()), 'last_month': str(latest.date()),
            'table_rows': len(meta), 'unique_price_codes': meta['Price code'].nunique(),
            'months': len(prices), 'price_observations': int(prices.notna().sum().sum()), 'tiers': tier_audit,
            'hierarchy_parent_checks': len(audit), 'workbook': output.name}
(ROOT / 'manifest.json').write_text(json.dumps(manifest, indent=2), encoding='utf-8')
book.close()
print(json.dumps(manifest, indent=2))
print(output)
