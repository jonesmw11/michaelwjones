# Download ABS national CPI tables and make a workbook with complete basket tiers.
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import pandas as pd
import requests
import xlsxwriter


ROOT = Path(__file__).resolve().parent
RAW = ROOT / 'raw'
RAW.mkdir(exist_ok=True)
BASE = 'https://www.abs.gov.au/statistics/economy/price-indexes-and-inflation/consumer-price-index-australia/jul-2026/'
WEIGHTS = 'https://www.abs.gov.au/system/files/83cf3a8e846b6b3f91ae7ea99ef968ab/Consumer%20Price%20Index%20-%202025%20Weighting%20Pattern.xlsx'
sources = {f'64010{i}.xlsx': BASE+f'64010{i}.xlsx' for i in list(range(1,10))+[17,18]}
sources['weights_2025.xlsx'] = WEIGHTS
manifest = []
for filename, url in sources.items():
    path = RAW / filename
    if not path.exists():
        response = requests.get(url, timeout=60)
        response.raise_for_status()
        path.write_bytes(response.content)
    manifest.append(dict(filename=filename,url=url,sha256=hashlib.sha256(path.read_bytes()).hexdigest(),retrieved_utc=datetime.now(timezone.utc).isoformat()))
    print('Source ready:', filename, flush=True)

# ABS publishes rounded weights; retain these and recursively reconcile sibling shares.
weights = openpyxl.load_workbook(RAW/'weights_2025.xlsx',data_only=True)['Table 1']
nodes = {'0':dict(code='0',name='All groups CPI',depth=0,parent=None,published=100.,children=[],weight=1.)}
stack = ['0']
for row in weights.iter_rows(min_row=8,max_col=6,values_only=True):
    if row[0] == 'ALL GROUPS CPI':
        break
    for col in range(3):
        if row[col] is not None:
            name = row[col]
            if col == 0:
                name = name.capitalize()
            depth = col+1
            parent = stack[depth-1]
            code = str(len(nodes))
            nodes[code]=dict(code=code,name=name,depth=depth,parent=parent,published=float(row[col+3]),children=[])
            nodes[parent]['children'].append(code)
            stack=stack[:depth]+[code]
for node in nodes.values():
    if node['children']:
        denominator = sum(nodes[c]['published'] for c in node['children'])
        for child in node['children']:
            nodes[child]['weight'] = node['weight']*nodes[child]['published']/denominator
assert [sum(n['depth']==i for n in nodes.values()) for i in range(1,4)] == [11,33,87]

records=[]
for filename in sources:
    if filename=='weights_2025.xlsx':
        continue
    workbook = openpyxl.load_workbook(RAW/filename,data_only=True,read_only=True)
    for sheet in workbook:
        if not sheet.title.startswith('Data'):
            continue
        rows=list(sheet.values)
        for col in range(1,len(rows[0])):
            if not rows[0][col]:
                continue
            description=rows[0][col]
            parts=[p.strip() for p in description.split(';') if p.strip()]
            frequency=rows[4][col]
            values={}
            for row in rows[10:]:
                if isinstance(row[0],datetime) and isinstance(row[col],(float,int)):
                    period=str(pd.Period(row[0],freq='Q' if frequency=='Quarter' else 'M'))
                    values[period]=row[col]
            records.append(dict(code=rows[9][col],description=description,measure=parts[0],name=parts[1] if len(parts)>1 else '',
                                frequency=frequency,unit=rows[1][col],adjustment=rows[2][col],source=filename,values=values))
    workbook.close()
periods=sorted({p for r in records for p in r['values']})
node_order=list(nodes)
positions={}
tier_records=[r for r in records if r['source'] in ('640103.xlsx','6401018.xlsx')]
for r in tier_records:
    key=(r['source'],r['measure'])
    position=positions.get(key,0)
    code=node_order[position]
    assert r['name'].casefold()==nodes[code]['name'].casefold(),(r['name'],nodes[code]['name'])
    r['node']=code
    positions[key]=position+1
panels={}
for r in tier_records:
    label=r['frequency']+' | '+r['measure']
    panels.setdefault(label,{})[r['node']]=r
panel_periods={label:sorted({p for r in items.values() for p in r['values']}) for label,items in panels.items()}
out=ROOT.parent/'Australia CPI - All Historical Data.xlsx'
checks=[]
with xlsxwriter.Workbook(out,{'constant_memory':True,'strings_to_formulas':False,'strings_to_urls':False}) as book:
    head=book.add_format({'bold':True,'bg_color':'#16624C','font_color':'white','text_wrap':True})
    wf=book.add_format({'num_format':'0.0000000000'})
    nf=book.add_format({'num_format':'0.00'})
    guide=book.add_worksheet('Guide and Weight Checks')
    lines=[
        'Australia CPI | Official ABS national history and complete basket tiers',
        'National coverage is the weighted average of eight capital cities; this is the ABS national CPI concept.',
        'Monthly and quarterly series are separate. Latest release: July 2026; quarterly data through June 2026.',
        'Index reference period and methodological changes follow the ABS source tables; no artificial splicing or imputation.',
        'Original rounded 2025 weights are preserved. Their tier totals are 100.00%, 100.01%, and 99.96%.',
        'Reconciled shares are DERIVED: parent share times child published weight / sum of published sibling weights.',
        'This produces coherent parent-child shares and every complete tier sums to 1. It does not recover unpublished exact ABS weights.',
        '2025 weights remain in use in 2026; they are not historical time-varying basket weights.',
        'Each tier has one row per component and all detailed monthly/quarterly measures across columns.',
        'Master includes national tables 1-9 and 17-18; city-specific tables 10-16 are outside this national scope.',
        'Blank observations are unavailable; overlapping analytical series in the master are not part of the basket tiers.',
        'Source: '+BASE,
        'Weights: '+WEIGHTS,
    ]
    for row,line in enumerate(lines):guide.write(row,0,line)
    guide.set_column(0,0,45);guide.set_column(1,5,22)
    guide.write_row(15,0,['Tier','Components','Published weight sum (%)','Reconciled share sum','Deviation from 1','Check'],head)
    master=book.add_worksheet('All National CPI Data')
    mh=['Series ID','Description','Frequency','Unit','Adjustment','Source file','First period','Last period']
    master.write_row(0,0,mh+periods,head);master.freeze_panes(1,2)
    master.set_column(0,0,18);master.set_column(1,1,65);master.set_column(2,7,18);master.set_column(8,7+len(periods),12,nf)
    pcols={p:8+i for i,p in enumerate(periods)}
    for row,r in enumerate(records,1):
        master.write_row(row,0,[r['code'],r['description'],r['frequency'],r['unit'],r['adjustment'],r['source'],min(r['values'],default=''),max(r['values'],default='')])
        for p,v in r['values'].items():master.write_number(row,pcols[p],v)
    master.autofilter(0,0,len(records),7+len(periods))
    for tier in range(1,4):
        members=[n for n in nodes.values() if n['depth']==tier]
        share_sum=sum(n['weight'] for n in members)
        assert abs(share_sum-1)<1e-12
        sheet=book.add_worksheet(f'Tier {tier}')
        sheet.write(0,0,f'Tier {tier}: '+['','11 groups','33 subgroups','87 expenditure classes'][tier])
        sheet.write(1,0,'Basket shares sum to 1')
        sheet.write_formula(1,4,f'=SUM(E4:E{len(members)+3})',wf,share_sum)
        header=['Name','Parent','Depth','Published weight (%)','Reconciled basket share','Adjustment from published share','Within-parent share']
        history_headers=[f'{label} | {p}' for label in panels for p in panel_periods[label]]
        sheet.write_row(2,0,header+history_headers,head);sheet.set_row(2,50);sheet.freeze_panes(3,2)
        sheet.set_column(0,1,42);sheet.set_column(2,6,20);sheet.set_column(7,6+len(history_headers),16,nf)
        for row,n in enumerate(members,3):
            parent=nodes[n['parent']]
            sibling_sum=sum(nodes[c]['published'] for c in parent['children'])
            sheet.write_row(row,0,[n['name'],nodes[n['parent']]['name'],tier,n['published']])
            sheet.write_formula(row,4,f'=D{row+1}/{sibling_sum!r}*{parent["weight"]!r}',wf,n['weight'])
            sheet.write_formula(row,5,f'=E{row+1}-D{row+1}/100',wf,n['weight']-n['published']/100)
            sheet.write_formula(row,6,f'=D{row+1}/{sibling_sum!r}',wf,n['published']/sibling_sum)
            col=7
            for label,items in panels.items():
                r=items[n['code']]
                for period in panel_periods[label]:
                    if period in r['values']:sheet.write_number(row,col,r['values'][period])
                    col+=1
        sheet.autofilter(2,0,len(members)+2,6+len(history_headers))
        r=15+tier
        guide.write_row(r,0,[tier,len(members),sum(n['published'] for n in members)])
        guide.write_formula(r,3,f"='Tier {tier}'!E2",wf,share_sum)
        guide.write_formula(r,4,f'=D{r+1}-1',wf,share_sum-1)
        guide.write_formula(r,5,f'=IF(ABS(E{r+1})<1E-12,"PASS","FAIL")',None,'PASS')
        checks.append(dict(tier=tier,components=len(members),published_percent_sum=sum(n['published'] for n in members),reconciled_sum=share_sum))
    hs=book.add_worksheet('Hierarchy and Weights')
    hs.write_row(0,0,['Name','Depth','Parent','Published %','Reconciled basket share','Children published %','Published rounding difference','Children reconciled share','Reconciled difference'],head)
    for row,n in enumerate(nodes.values(),1):
        childsum=sum(nodes[c]['published'] for c in n['children']) if n['children'] else None
        shares=sum(nodes[c]['weight'] for c in n['children']) if n['children'] else None
        if shares is not None:assert abs(shares-n['weight'])<1e-12
        hs.write_row(row,0,[n['name'],n['depth'],n['parent'],n['published'],n['weight'],childsum,childsum-n['published'] if childsum is not None else None,shares,shares-n['weight'] if shares is not None else None])
    hs.set_column(0,0,45);hs.set_column(1,8,24);hs.freeze_panes(1,1)
    guide.activate()
(ROOT/'manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
(ROOT/'weight_checks.json').write_text(json.dumps(checks,indent=2),encoding='utf-8')
pd.DataFrame([{k:v for k,v in n.items() if k!='children'} for n in nodes.values()]).to_csv(ROOT/'hierarchy_weights.csv',index=False)
pd.DataFrame([{k:v for k,v in r.items() if k not in ('values','node')} for r in records]).to_csv(ROOT/'series_catalog.csv',index=False)
print(out,flush=True)
print(json.dumps(checks),flush=True)
print('Master series:',len(records),'observations:',sum(len(r['values']) for r in records),'period range:',min(periods),max(periods),flush=True)
