# Assemble published core/underlying measures, preserving units, seasonal adjustment and source definitions.
import hashlib
import json
from datetime import datetime
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parent
BASE = ROOT.parent
JP = BASE / 'JP Inflation Model' / 'japan_official'
RAW = ROOT / 'raw'
catalog, series = [], {}
downloads = {r['file']: r['url'] for r in json.loads((ROOT / 'download_manifest.json').read_text())}


def add(country, name, values, unit, frequency, adjustment, source, note='', kind='Published'):
    values = pd.to_numeric(values, errors='coerce').dropna().sort_index()
    values.index = values.index.astype(str)
    assert values.index.is_unique and np.isfinite(values).all()
    if values.empty:
        return
    sid = f'{country[:2].upper()}{len(catalog)+1:04d}'
    series[sid] = values
    catalog.append(dict(id=sid, country=country, name=name, unit=unit, frequency=frequency, adjustment=adjustment,
                        source=source, kind=kind, notes=note, first_period=values.index.min(), last_period=values.index.max(), observations=len(values)))


def add_index(country, name, values, unit, frequency, adjustment, source, note=''):
    add(country, name, values, unit, frequency, adjustment, source, note)
    # Reindex the calendar first so gaps cannot silently turn into multi-period changes.
    freq = 'M' if frequency == 'Monthly' else 'Q'
    v = pd.to_numeric(values, errors='coerce').copy()
    v.index = pd.PeriodIndex(v.index, freq=freq)
    v = v.reindex(pd.period_range(v.index.min(), v.index.max(), freq=freq))
    for lag, label in [(1, 'MoM percent' if freq == 'M' else 'QoQ percent'), (12 if freq == 'M' else 4, 'YoY percent')]:
        rate = v.pct_change(lag, fill_method=None) * 100
        add(country, name, rate, label, frequency, adjustment, source, 'Derived from published index; no filling. ' + note, 'Derived')


# Korea: official rendered tables; two bounded extracts overlap in seven months.
def korea_read(filename, expected_sum=None):
    obj = json.loads((RAW / filename).read_text())
    dates = pd.period_range(obj['start'], obj['end'], freq='M').astype(str)
    assert len(dates) == len(obj['values'])
    if expected_sum is not None:
        assert abs(sum(obj['values']) - expected_sum) < 1e-7
    return obj, pd.Series(obj['values'], index=dates)


kf, food = korea_read('korea_food_energy.json', 35316.095)
ka, early = korea_read('korea_ag_oil_early.json', 26804.105)
_, recent = korea_read('korea_ag_oil_recent.json')
for data, expected in [(food, 8906493.844), (early, 8535422.822), (recent, 896008.808)]:
    assert abs(float(np.dot(data.values, np.arange(1, len(data)+1))) - expected) < 1e-6
overlap = early.index.intersection(recent.index)
assert len(overlap) == 7
np.testing.assert_array_equal(early[overlap], recent[overlap])
ag = recent.combine_first(early).sort_index()
assert len(ag) == 620 and len(food) == 440
for name, vals, source in [('CPI excluding food and energy', food, kf['source']),
                           ('CPI excluding agricultural products and oils', ag, ka['source'])]:
    add_index('Korea', name, vals, 'Index 2020=100', 'Monthly', 'Not seasonally adjusted', source,
              'Statistics Korea/MODS via KOSIS. Retrieved from the public rendered table; not dependent on the incomplete detailed Korea workbook.')

# Japan: all national exclusion aggregates from the existing official e-Stat tables.
for entry in json.loads((JP / 'manifest.json').read_text()):
    if entry['frequency'] != 'monthly':
        continue
    stem = f'monthly_{entry["file_id"]}'
    metadata = pd.read_csv(JP / 'metadata' / f'{stem}.csv', dtype=str).set_index('series_code')
    frame = pd.read_csv(JP / 'processed' / f'{stem}.csv', index_col=0)
    title = entry['title']
    unit = 'MoM percent' if 'previous month' in title else 'YoY percent' if 'over the year' in title else 'Index 2025=100'
    adjustment = 'Seasonally adjusted' if 'Seasonally Adjusted' in title else 'Not seasonally adjusted'
    for code, row in metadata.iterrows():
        if str(row['name_en']).startswith('All items, less'):
            add('Japan', row['name_en'], frame[code], unit, 'Monthly', adjustment, entry['url'], title)

# BOJ's base-specific research histories are kept separate; comparable rate series are also joined latest-base-first.
wb = openpyxl.load_workbook(RAW / 'japan_boj.xlsx', read_only=True, data_only=True)
rows = list(wb['chart'].values)
groups = {}
for col in range(1, 53):
    name, vintage = rows[3][col], rows[4][col]
    if not name or not vintage:
        continue
    values = pd.Series({r[0].strftime('%Y-%m'): r[col] for r in rows[5:] if isinstance(r[0], datetime)})
    unit = 'YoY percent' if 'y/y' in name else 'Percentage points' if 'Diffusion' in name else 'Percent of items'
    if unit != 'YoY percent':
        continue
    add('Japan', 'BOJ ' + name + ' | ' + vintage, values, unit, 'Monthly', 'As published', downloads['japan_boj.xlsx'],
        'Base-specific BOJ research estimate; institutional-factor adjustments differ from the standard Statistics Bureau core.')
    groups.setdefault(name, []).append((vintage, pd.to_numeric(values, errors='coerce')))
for name, parts in groups.items():
    combined = pd.Series(dtype=float)
    for vintage, values in sorted(parts, reverse=True):
        combined = combined.combine_first(values)
    add('Japan', 'BOJ ' + name + ' | joined history', combined, 'YoY percent', 'Monthly', 'As published', downloads['japan_boj.xlsx'],
        'Joins source rate columns, prioritising newest base at overlaps; base definitions change. Raw base-specific series are also retained.', 'Joined published rates')
for col in range(55, 60):
    values = {}
    for r in rows[5:]:
        if isinstance(r[54], (int, float)):
            year = int(r[54]); quarter = round((r[54] - year) * 10)
            if 1 <= quarter <= 4:
                values[f'{year}Q{quarter}'] = r[col]
    add('Japan', 'BOJ ' + rows[3][col], pd.Series(values), 'YoY percent', 'Quarterly', 'As published', downloads['japan_boj.xlsx'],
        'Quarterly figures as published; the current quarter may contain only the months released so far.')
wb.close()

# Australia: ABS analytical exclusions and distribution-based cores, at each published frequency/adjustment.
au = BASE / 'AU Inflation Model' / 'australia_official'
for path in (au / 'raw').glob('64010*.xlsx'):
    if 'Appendix' in path.name:
        continue
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for sheet in wb:
        if not sheet.title.startswith('Data'):
            continue
        rows = list(sheet.values)
        for col in range(1, len(rows[0])):
            description = rows[0][col]
            if not description or not any(word in str(description).lower() for word in ['excluding', 'trimmed mean', 'weighted median']):
                continue
            if 'Australia' not in description and 'Weighted average' not in description:
                continue
            frequency = 'Quarterly' if rows[4][col] == 'Quarter' else 'Monthly'
            freq = 'Q' if frequency == 'Quarterly' else 'M'
            vals = pd.Series({str(pd.Period(r[0], freq=freq)): r[col] for r in rows[10:] if isinstance(r[0], datetime)})
            parts = [p.strip() for p in description.split(';') if p.strip()]
            add('Australia', parts[1], vals, parts[0], frequency, str(rows[2][col]),
                'https://www.abs.gov.au/statistics/economy/price-indexes-and-inflation/consumer-price-index-australia/jul-2026/' + path.name,
                'ABS analytical series. Exclusion variants are not all headline core measures. Source series ID: ' + str(rows[9][col]))
    wb.close()
wb = openpyxl.load_workbook(RAW / 'australia_rba.xlsx', read_only=True, data_only=True)
rows = list(wb['Data'].values)
for col in range(3, 21):
    vals = pd.Series({str(pd.Period(r[0], freq='Q')): r[col] for r in rows[11:] if isinstance(r[0], datetime)})
    add('Australia', 'RBA ' + rows[1][col], vals, 'YoY percent' if col < 11 else 'QoQ percent', 'Quarterly', rows[4][col], downloads['australia_rba.xlsx'],
        str(rows[2][col]) + '. Quarterly-collection basis, see original workbook Notes. ID: ' + str(rows[10][col]))
wb.close()

# USA: all aggregate exclusion variants already available in the CPI and PCE workbooks.
us = BASE / 'USA Inflation Model' / 'usa_official'
cat = pd.read_csv(us / 'series_catalog.csv')
frame = pd.read_csv(us / 'national_cpi_indices.csv', index_col=0)
for _, row in cat.iterrows():
    name = str(row['name']).strip()
    if 'less' in name.lower() and name.startswith(('All items', 'Services', 'Commodities')):
        add_index('USA', 'CPI: ' + name, frame[row['code']], 'Native CPI index base ' + str(row['base_code']), 'Monthly', 'Not seasonally adjusted',
                  'https://www.bls.gov/cpi/data.htm', 'BLS data; pre-2017 mostly DBnomics mirror, 2017 onward directly BLS. ' + row['code'])
pce = BASE / 'USA Inflation Model' / 'pce_official'
cat = pd.read_csv(pce / 'series.csv')
frame = pd.read_csv(pce / 'monthly_prices.csv', index_col=0)
for _, row in cat.iterrows():
    if 'excluding' in row['Name'].lower():
        add_index('USA', row['Name'], frame[str(row['Line'])], 'Index 2017=100', 'Monthly', 'Seasonally adjusted',
                  'https://apps.bea.gov/national/Release/XLS/Underlying/Section2All_xls.xlsx', 'BEA Table 2.4.4U; ' + row['Price code'])

# Cleveland rates include raw decimal monthly changes, explicitly converted to percentage units.
for filename, name in [('usa_median_cpi.csv', 'Median CPI'), ('usa_trimmed_cpi.csv', '16% trimmed-mean CPI')]:
    df = pd.read_csv(RAW / filename)
    df.index = pd.to_datetime(df.iloc[:, 0]).dt.to_period('M').astype(str)
    monthly = pd.to_numeric(df.iloc[:, 1], errors='coerce') * 100
    add('USA', 'Cleveland ' + name, monthly, 'MoM percent', 'Monthly', 'Seasonally adjusted', downloads[filename], 'Source decimal monthly changes multiplied by 100.')
    add('USA', 'Cleveland ' + name, df.iloc[:, 2], '1-month annualised percent', 'Monthly', 'Seasonally adjusted', downloads[filename])
    monthly.index = pd.PeriodIndex(monthly.index, freq='M')
    monthly = monthly.reindex(pd.period_range(monthly.index.min(), monthly.index.max(), freq='M'))
    yoy = ((1 + monthly / 100).rolling(12, min_periods=12).apply(np.prod, raw=True) - 1) * 100
    add('USA', 'Cleveland ' + name, yoy, 'YoY percent', 'Monthly', 'Seasonally adjusted', downloads[filename], 'Compounded twelve source monthly rates; no bridging gaps.', 'Derived')
df = pd.read_csv(RAW / 'usa_core_cpi.csv')
df.index = pd.to_datetime(df.iloc[:, 0]).dt.to_period('M').astype(str)
add_index('USA', 'Core CPI excluding food and energy (Cleveland/BLS)', df['core_index'], 'Index 1982-84=100', 'Monthly', 'Seasonally adjusted', downloads['usa_core_cpi.csv'])
df = pd.read_csv(RAW / 'usa_median_pce.csv')
df.index = pd.to_datetime(df.iloc[:, 0]).dt.to_period('M').astype(str)
for col, unit in [(1, 'MoM percent'), (2, 'YoY percent')]:
    add('USA', 'Cleveland median PCE', df.iloc[:, col], unit, 'Monthly', 'Seasonally adjusted', downloads['usa_median_pce.csv'])
wb = openpyxl.load_workbook(RAW / 'usa_dallas_history.xlsx', read_only=True, data_only=True)
rows = list(wb.worksheets[0].values)
for col, unit in [(1, '1-month annualised percent'), (2, '6-month annualised percent'), (3, 'YoY percent')]:
    vals = pd.Series({r[0].strftime('%Y-%m'): r[col] for r in rows[4:] if isinstance(r[0], datetime)})
    add('USA', 'Dallas trimmed-mean PCE', vals, unit, 'Monthly', 'Seasonally adjusted', downloads['usa_dallas_history.xlsx'])
wb.close()
wb = openpyxl.load_workbook(RAW / 'usa_sticky.xlsx', read_only=True, data_only=True)
rows = list(wb['Data'].values)
for start in range(1, 25, 4):
    for offset, unit in enumerate(['MoM percent', '1-month annualised percent', '3-month annualised percent', 'YoY percent']):
        vals = pd.Series({r[0].strftime('%Y-%m'): r[start+offset] for r in rows[1:] if isinstance(r[0], datetime)})
        vals = pd.to_numeric(vals, errors='coerce') * (100 if offset == 0 else 1)
        add('USA', 'Atlanta ' + rows[0][start], vals, unit, 'Monthly', 'Seasonally adjusted', downloads['usa_sticky.xlsx'],
            'Sticky/flexible and exclusion variants retained. Monthly decimals converted to percent.')
wb.close()
df = pd.read_csv(RAW / 'usa_mct.csv', skiprows=3)
df.index = pd.to_datetime(df.iloc[:, 1]).dt.to_period('M').astype(str)
for col, suffix in [(2, 'lower band'), (3, 'point estimate'), (4, 'upper band')]:
    add('USA', 'NY Fed Multivariate Core Trend | ' + suffix, df.iloc[:, col], 'Annual inflation percent (model trend)', 'Monthly', 'Model estimate', downloads['usa_mct.csv'],
        'PCE-based dynamic-factor trend. These are model estimates, not realised YoY changes. Source chart parser maps CSV columns 2/3/4 to lower/point/upper.')

metadata = pd.DataFrame(catalog).set_index('id')
notes = [
    'Core and underlying inflation | USA, Japan, Australia, Korea',
    'The catalogue defines coverage. Includes standard national cores, published exclusion variants, and regularly released central-bank underlying measures.',
    'This is not an assertion that every experimental or discontinued research measure worldwide is included.',
    'Country sheets retain each series in its native units; consult Series for definition, adjustment, frequency, source and coverage.',
    'Monthly and quarterly observations are separate. Rates are percent (2 means 2%). Index bases are series-specific.',
    'Cores overlap. Do not sum these measures or their weights as though they form a basket.',
    'Derived rates use exact calendar lags with no filling. Annualised rates remain explicitly labelled.',
    'Japan BOJ base-specific estimates and joined rate histories are both retained. Joined histories prefer the newest base at overlaps.',
    'Korea: complete core aggregate histories retrieved from KOSIS; this does NOT complete its outstanding detailed component workbook.',
    'Korea excluding agricultural products and oils: 1975-01 to 2026-08. Excluding food and energy: 1990-01 to 2026-08.',
    'USA CPI mirror provenance remains documented; Fed, BEA, BOJ, RBA and ABS additions come directly from their publishers.',
    'Source files and scripts are retained in core_inflation. Most downloads are cached; Korea currently requires browser table extraction for updates.'
]


def write_workbook(path, countries):
    with pd.ExcelWriter(path, engine='xlsxwriter', engine_kwargs={'options': {'strings_to_formulas': False}}) as writer:
        pd.DataFrame({'Guide': notes}).to_excel(writer, sheet_name='Guide', index=False)
        subset = metadata[metadata['country'].isin(countries)]
        subset.to_excel(writer, sheet_name='Series')
        for country in countries:
            for frequency, short in [('Monthly', 'M'), ('Quarterly', 'Q')]:
                ids = subset.index[(subset.country == country) & (subset.frequency == frequency)]
                if len(ids):
                    panel = pd.DataFrame({sid: series[sid] for sid in ids}).sort_index()
                    panel.columns = [f'{sid} | {metadata.loc[sid, "name"]} | {metadata.loc[sid, "unit"]} | {metadata.loc[sid, "adjustment"]}' for sid in ids]
                    panel.index.name = 'Period'
                    panel.to_excel(writer, sheet_name=f'{country} {short}')
        for name, sheet in writer.sheets.items():
            sheet.freeze_panes(1, 1)
            sheet.set_column(0, 0, 18)
            sheet.set_column(1, 500, 17)
            if name != 'Guide':
                sheet.autofilter(0, 0, sheet.dim_rowmax, sheet.dim_colmax)
                sheet.set_row(0, 70)
        writer.sheets['Guide'].set_column(0, 0, 145)
        writer.sheets['Series'].set_column(2, 2, 75)
        writer.sheets['Series'].set_column(7, 9, 60)


output = BASE / 'Core Inflation - All Countries.xlsx'
write_workbook(output, ['USA', 'Japan', 'Australia', 'Korea'])
country_directories = {
    'USA': BASE / 'USA Inflation Model',
    'Japan': BASE / 'JP Inflation Model',
    'Australia': BASE / 'AU Inflation Model',
    'Korea': BASE / 'KR Inflation Model',
}
for country in ['USA', 'Japan', 'Australia', 'Korea']:
    directory = country_directories[country]
    write_workbook(directory / f'{country} - Core Inflation Measures.xlsx', [country])
metadata.to_csv(ROOT / 'series_catalog.csv', encoding='utf-8-sig')
long = pd.concat([s.rename('Value').rename_axis('Period').reset_index().assign(id=sid) for sid, s in series.items()], ignore_index=True)
long.to_csv(ROOT / 'observations.csv', index=False)
report = {'countries': metadata.groupby('country').size().to_dict(), 'observations': len(long),
          'korea_food_energy_count': len(food), 'korea_ag_oil_count': len(ag), 'korea_overlap_matches': len(overlap),
          'raw_hashes': {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in RAW.iterdir() if p.is_file()}}
(ROOT / 'build_checks.json').write_text(json.dumps(report, indent=2))
print(output)
print(json.dumps({k: v for k, v in report.items() if k != 'raw_hashes'}, indent=2))
